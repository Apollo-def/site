import io
from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from auditoria.models import LogAuditoria
from contas.importacao import gerar_modelo_planilha, importar_funcionarios
from contas.models import Contrato, Funcionario

TITULOS = ["Nome completo", "Matrícula", "CPF", "Data de nascimento", "Contrato", "Cargo", "Situação"]


def planilha_xlsx(linhas, titulos=TITULOS, linhas_antes=0):
    livro = Workbook()
    aba = livro.active
    for _ in range(linhas_antes):
        aba.append(["GRUPO GLOBAL - RELAÇÃO DE FUNCIONÁRIOS"])
    aba.append(titulos)
    for linha in linhas:
        aba.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    saida.seek(0)
    return saida


class ImportacaoTestes(TestCase):
    def setUp(self):
        self.semed = Contrato.objects.create(sigla="SEMED")
        Contrato.objects.create(sigla="TRT-16")

    def test_cria_funcionarios_sem_login(self):
        arquivo = planilha_xlsx([
            ["Maria da Silva", "1001", "529.982.247-25", datetime(1990, 3, 7), "semed", "Porteira", "Ativo"],
            ["João Souza", 1002, 11144477735, "15/08/1985", "TRT 16", "Auxiliar de limpeza", ""],
        ])
        r = importar_funcionarios(arquivo, "func.xlsx")
        self.assertEqual((len(r.criados), len(r.erros)), (2, 0), r.erros)
        joao = Funcionario.objects.get(matricula="1002")
        self.assertEqual(joao.cpf, "11144477735")
        self.assertEqual(joao.data_nascimento, date(1985, 8, 15))
        self.assertEqual(joao.contrato.sigla, "TRT-16")
        self.assertIsNone(joao.usuario)  # atendimento pelo WhatsApp, sem login

    def test_erros_por_linha_nao_impedem_as_outras(self):
        arquivo = planilha_xlsx([
            ["Maria", "1001", "529.982.247-25", "07/03/1990", "SEMED", "Porteira", ""],
            ["CPF errado", "1002", "529.982.247-26", "07/03/1990", "SEMED", "Porteira", ""],
            ["Repetida", "1001", "111.444.777-35", "07/03/1990", "SEMED", "Porteira", ""],
            ["Contrato novo", "1003", "111.444.777-35", "07/03/1990", "XYZ", "Porteira", ""],
            ["Data ruim", "1004", "012.345.678-90", "31/02/1990", "SEMED", "Porteira", ""],
            [None, None, None, None, None, None, None],  # linha vazia é ignorada
        ])
        r = importar_funcionarios(arquivo, "func.xlsx")
        self.assertEqual(len(r.criados), 1)
        self.assertEqual([linha for linha, _, _ in r.erros], [3, 4, 5, 6])
        self.assertIn("CPF inválido", r.erros[0][2])
        self.assertIn("repetida", r.erros[1][2])
        self.assertIn("não está cadastrado", r.erros[2][2])
        self.assertIn("inválida", r.erros[3][2])

    def test_atualiza_por_matricula_e_detecta_sem_alteracao(self):
        linha = ["Maria", "1001", "52998224725", "07/03/1990", "SEMED", "Porteira", ""]
        importar_funcionarios(planilha_xlsx([linha]), "a.xlsx")
        r = importar_funcionarios(planilha_xlsx([linha]), "a.xlsx")
        self.assertEqual(r.sem_alteracao, 1)
        r = importar_funcionarios(planilha_xlsx([linha[:5] + ["Recepcionista", "Inativo"]]), "a.xlsx")
        self.assertEqual(len(r.atualizados), 1)
        self.assertIn("cargo", r.atualizados[0][2])
        maria = Funcionario.objects.get(matricula="1001")
        self.assertEqual(maria.situacao, Funcionario.Situacao.INATIVO)
        self.assertEqual(Funcionario.objects.count(), 1)

    def test_cpf_de_outro_funcionario_e_recusado(self):
        importar_funcionarios(planilha_xlsx([["Maria", "1001", "52998224725", "07/03/1990", "SEMED", "x", ""]]), "a.xlsx")
        r = importar_funcionarios(planilha_xlsx([["Outra", "2002", "52998224725", "07/03/1990", "SEMED", "x", ""]]), "a.xlsx")
        self.assertIn("outro funcionário", r.erros[0][2])

    def test_simulacao_nao_salva_nada(self):
        arquivo = planilha_xlsx([["Maria", "1001", "52998224725", "07/03/1990", "NOVO", "x", ""]])
        r = importar_funcionarios(arquivo, "a.xlsx", criar_contratos=True, simulacao=True)
        self.assertEqual(len(r.criados), 1)
        self.assertEqual(r.contratos_criados, ["NOVO"])
        self.assertFalse(Funcionario.objects.exists())
        self.assertFalse(Contrato.objects.filter(sigla="NOVO").exists())
        self.assertFalse(get_user_model().objects.exists())

    def test_titulos_alternativos_linha_de_titulo_e_zeros_perdidos(self):
        titulos = ["MAT", "NOME", "CPF", "DT NASC", "LOTAÇÃO", "FUNÇÃO"]
        arquivo = planilha_xlsx([["77", "Ana", 1234567890, "01/01/1980", "SEMED", "Copeira"]], titulos, linhas_antes=2)
        r = importar_funcionarios(arquivo, "a.xlsx")
        self.assertEqual(len(r.criados), 1, r.erros)
        self.assertEqual(Funcionario.objects.get().cpf, "01234567890")

    def test_colunas_obrigatorias_faltando(self):
        r = importar_funcionarios(planilha_xlsx([], ["Matrícula", "CPF", "Nome"]), "a.xlsx")
        self.assertIn("Data de nascimento", r.erro_geral)

    def test_csv_do_excel_em_portugues(self):
        conteudo = "Nome completo;Matrícula;CPF;Data de nascimento;Contrato;Cargo\n" \
                   "José Conceição;1001;529.982.247-25;07/03/1990;SEMED;Porteiro\n"
        r = importar_funcionarios(io.BytesIO(conteudo.encode("cp1252")), "a.csv")
        self.assertEqual(len(r.criados), 1, r.erro_geral or r.erros)
        self.assertEqual(Funcionario.objects.get().nome_completo, "José Conceição")

    def test_modelo_de_planilha_e_reconhecido(self):
        livro = load_workbook(io.BytesIO(gerar_modelo_planilha()))
        self.assertEqual(livro.sheetnames, ["Funcionários", "Instruções"])
        r = importar_funcionarios(io.BytesIO(gerar_modelo_planilha()), "modelo.xlsx")
        self.assertEqual((r.erro_geral, r.total_linhas), ("", 0))


class ImportacaoTelaTestes(TestCase):
    def setUp(self):
        Contrato.objects.create(sigla="SEMED")
        self.client.force_login(get_user_model().objects.create_superuser("rh", "rh@x.com", "SenhaForte!2026"))

    def test_tela_importa_e_registra_auditoria(self):
        url = reverse("admin:contas_funcionario_importar")
        self.assertEqual(self.client.get(url).status_code, 200)
        arquivo = SimpleUploadedFile("f.xlsx", planilha_xlsx(
            [["Maria", "1001", "52998224725", "07/03/1990", "SEMED", "x", ""]]).getvalue())
        resposta = self.client.post(url, {"arquivo": arquivo})
        self.assertContains(resposta, "Importação concluída sem erros")
        self.assertTrue(Funcionario.objects.filter(matricula="1001").exists())
        self.assertTrue(LogAuditoria.objects.filter(acao=LogAuditoria.Acao.IMPORTACAO_FUNCIONARIOS).exists())

    def test_download_do_modelo(self):
        resposta = self.client.get(reverse("admin:contas_funcionario_modelo_planilha"))
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("modelo_importacao_funcionarios.xlsx", resposta["Content-Disposition"])

    def test_botoes_aparecem_na_lista(self):
        resposta = self.client.get(reverse("admin:contas_funcionario_changelist"))
        self.assertContains(resposta, "Importar planilha")


class PlanilhaSimplesTestes(TestCase):
    """Planilha só com nome, matrícula, CPF e data de nascimento."""

    TITULOS = ["NOME", "MATRÍCULA", "CPF", "DATA DE NASCIMENTO"]

    def setUp(self):
        self.semed = Contrato.objects.create(sigla="SEMED")

    def test_usa_contrato_padrao(self):
        arquivo = planilha_xlsx([["Maria", "1001", "529.982.247-25", "07/03/1990"]], self.TITULOS)
        r = importar_funcionarios(arquivo, "a.xlsx", contrato_padrao=self.semed)
        self.assertEqual((len(r.criados), r.erros), (1, []))
        maria = Funcionario.objects.get()
        self.assertEqual((maria.contrato, maria.cargo), (self.semed, ""))

    def test_sem_contrato_padrao_avisa(self):
        r = importar_funcionarios(planilha_xlsx([["Maria", "1001", "52998224725", "07/03/1990"]], self.TITULOS), "a.xlsx")
        self.assertIn("contrato padrão", r.erros[0][2])

    def test_reimportar_mantem_contrato_e_cargo(self):
        tjma = Contrato.objects.create(sigla="TJMA")
        Funcionario.objects.create(nome_completo="Maria", matricula="1001", cpf="52998224725",
                                   data_nascimento=date(1990, 3, 7), contrato=tjma, cargo="Porteira")
        r = importar_funcionarios(planilha_xlsx([["Maria", "1001", "52998224725", "07/03/1990"]], self.TITULOS),
                                  "a.xlsx", contrato_padrao=self.semed)
        self.assertEqual(r.sem_alteracao, 1)
        maria = Funcionario.objects.get()
        self.assertEqual((maria.contrato, maria.cargo), (tjma, "Porteira"))


class DepartamentoComoContratoTestes(TestCase):
    def test_departamento_que_comeca_com_a_sigla(self):
        Contrato.objects.create(sigla="SEMED")
        Contrato.objects.create(sigla="SEMCAS")
        titulos = ["Cód Epr", "Nome", "CPF", "Data nascimento", "Descrição Dpto"]
        arquivo = planilha_xlsx([
            ["4830", "Maria", "52998224725", "07/03/1990", "SEMED - CT 11/2024 FERISTA"],
            ["6168", "João", "11144477735", "01/01/1980", "SEMCAS CT 43/2024"],
            ["7000", "Ana", "01234567890", "01/01/1980", "OUTRO ÓRGÃO"],
        ], titulos)
        r = importar_funcionarios(arquivo, "a.xlsx")
        self.assertEqual(len(r.criados), 2)
        self.assertEqual(Funcionario.objects.get(matricula="4830").contrato.sigla, "SEMED")
        self.assertEqual(Funcionario.objects.get(matricula="6168").contrato.sigla, "SEMCAS")
        self.assertEqual(r.contratos_associados["SEMED - CT 11/2024 FERISTA"], "SEMED")
        self.assertIn("não está cadastrado", r.erros[0][2])
