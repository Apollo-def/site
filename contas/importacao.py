"""
Importação de funcionários por planilha (Excel .xlsx ou CSV).

Regras:
- O funcionário é identificado pela MATRÍCULA. Se já existir, os dados são atualizados;
  se não existir, é cadastrado.
- Os títulos das colunas podem variar ("Matrícula", "MAT", "Lotação", "Dt Nasc"...): veja COLUNAS.
- Cada linha é independente: uma linha com erro não impede as outras de entrarem.
- No modo "somente conferir", tudo é processado e depois desfeito (nada fica salvo).
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Contrato, Funcionario
from .validadores import extrair_celular, formatar_celular, somente_digitos

# campo do sistema: (título no modelo de planilha, obrigatório?, outros títulos aceitos)
COLUNAS = {
    "nome_completo": ("Nome completo", True, ["nome", "nomedofuncionario", "funcionario", "colaborador",
                                              "nomedocolaborador", "empregado"]),
    "matricula": ("Matrícula", True, ["mat", "matr", "matriculafuncionario", "registro", "nmatricula", "codepr",
                                      "codempregado", "codfuncionario", "codigoempregado", "codigofuncionario"]),
    "cpf": ("CPF", True, ["ncpf", "cpfdofuncionario"]),
    "data_nascimento": ("Data de nascimento", True, ["datanascimento", "nascimento", "dtnascimento", "dtnasc",
                                                     "datanasc", "dnasc"]),
    "contrato": ("Contrato", False, ["lotacao", "contratolotacao", "orgao", "cliente", "tomador", "descricaodpto",
                                     "dpto", "departamento", "descricaodepartamento"]),
    "cargo": ("Cargo", False, ["funcao", "cargofuncao", "tipodeposto", "tipoposto"]),
    "instalacao": ("Instalação", False, ["instalacoes", "localdetrabalho", "unidade"]),
    "data_admissao": ("Data de admissão", False, ["dataadmissao", "admissao", "dtadmissao", "dtadm"]),
    "telefone": ("Telefone", False, ["celular", "whatsapp", "fone", "fones", "telefonewhatsapp", "contato", "tel"]),
    "situacao": ("Situação", False, ["status", "sit"]),
    "data_desligamento": ("Data de desligamento", False, ["datadesligamento", "desligamento", "dtdesligamento",
                                                          "datademissao", "demissao"]),
}
CAMPOS_OBRIGATORIOS = [campo for campo, (_, obrigatorio, _) in COLUNAS.items() if obrigatorio]
LIMITE_LINHAS = 20000


class ErroPlanilha(Exception):
    """Problema no arquivo inteiro (formato, colunas faltando)."""


class ErroLinha(Exception):
    """Problema em uma linha específica."""


@dataclass
class ResultadoImportacao:
    simulacao: bool = False
    total_linhas: int = 0
    criados: list = field(default_factory=list)       # [(linha, "Nome (matrícula)")]
    atualizados: list = field(default_factory=list)   # [(linha, "Nome (matrícula)", "campos alterados")]
    sem_alteracao: int = 0
    erros: list = field(default_factory=list)         # [(linha, identificação, mensagem)]
    contratos_criados: list = field(default_factory=list)
    contratos_associados: dict = field(default_factory=dict)  # texto da planilha -> sigla do contrato
    colunas_reconhecidas: list = field(default_factory=list)
    erro_geral: str = ""

    def resumo(self):
        texto = (f"{len(self.criados)} criados, {len(self.atualizados)} atualizados, "
                 f"{self.sem_alteracao} sem alteração, {len(self.erros)} com erro")
        return texto + (" (simulação — nada foi salvo)" if self.simulacao else "")


# ---------------------------------------------------------------------------
# Leitura do arquivo
# ---------------------------------------------------------------------------
def normalizar(texto):
    """'Data de Nascimento' -> 'datadenascimento' (sem acento, sem espaço, minúsculo)."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def _titulos_aceitos():
    aceitos = {}
    for campo, (titulo, _, apelidos) in COLUNAS.items():
        for nome in [titulo, campo, *apelidos]:
            aceitos[normalizar(nome)] = campo
    return aceitos


def ler_planilha(arquivo, nome_arquivo):
    """Devolve uma lista de (número da linha, [valores das células])."""
    extensao = nome_arquivo.lower().rsplit(".", 1)[-1]
    if extensao == "xlsx":
        try:
            planilha = load_workbook(arquivo, read_only=True, data_only=True).worksheets[0]
        except Exception as erro:
            raise ErroPlanilha("Não foi possível abrir o arquivo Excel. Ele está corrompido ou protegido por senha?") from erro
        return [(numero, list(valores)) for numero, valores in enumerate(planilha.iter_rows(values_only=True), 1)]

    if extensao == "xls":
        import xlrd  # formato antigo do Excel (97-2003)
        try:
            livro = xlrd.open_workbook(file_contents=arquivo.read())
        except Exception as erro:
            raise ErroPlanilha("Não foi possível abrir o arquivo .xls. Abra no Excel e salve como .xlsx.") from erro
        aba = livro.sheet_by_index(0)

        def valor(linha, coluna):
            if aba.cell_type(linha, coluna) == xlrd.XL_CELL_DATE:
                return xlrd.xldate_as_datetime(aba.cell_value(linha, coluna), livro.datemode)
            return aba.cell_value(linha, coluna)

        return [(linha + 1, [valor(linha, c) for c in range(aba.ncols)]) for linha in range(aba.nrows)]

    if extensao == "csv":
        bruto = arquivo.read()
        try:
            conteudo = bruto.decode("utf-8-sig")
        except UnicodeDecodeError:
            conteudo = bruto.decode("latin-1")  # CSV salvo pelo Excel em português
        primeira_linha = conteudo.split("\n", 1)[0]
        separador = max([";", ",", "\t", "|"], key=primeira_linha.count)
        leitor = csv.reader(io.StringIO(conteudo), delimiter=separador)
        return [(numero, valores) for numero, valores in enumerate(leitor, 1)]

    raise ErroPlanilha("Formato não aceito. Envie um arquivo Excel (.xlsx ou .xls) ou .csv.")


def localizar_cabecalho(linhas):
    """Procura, nas primeiras linhas, a linha de títulos. Devolve (posição na lista, {coluna: campo})."""
    aceitos = _titulos_aceitos()
    for posicao, (_, valores) in enumerate(linhas[:15]):
        mapa = {}
        for coluna, valor in enumerate(valores):
            campo = aceitos.get(normalizar(valor))
            if campo and campo not in mapa.values():
                mapa[coluna] = campo
        if {"matricula", "cpf"} <= set(mapa.values()):
            faltando = [COLUNAS[c][0] for c in CAMPOS_OBRIGATORIOS if c not in mapa.values()]
            if faltando:
                raise ErroPlanilha("A planilha não tem as colunas obrigatórias: " + ", ".join(faltando) + ".")
            return posicao, mapa
    raise ErroPlanilha("Não encontrei a linha de títulos. A planilha precisa ter pelo menos as colunas "
                       + ", ".join(COLUNAS[c][0] for c in CAMPOS_OBRIGATORIOS) + ".")


# ---------------------------------------------------------------------------
# Conversão dos valores das células
# ---------------------------------------------------------------------------
def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # Excel transforma 1234 em 1234.0
    return " ".join(str(valor).split())


def converter_cpf(valor):
    cpf = somente_digitos(texto(valor))
    # O Excel costuma apagar os zeros à esquerda (012.345... vira 12345...).
    return cpf.zfill(11) if 0 < len(cpf) < 11 else cpf


def converter_data(valor, nome_campo):
    if valor is None or texto(valor) == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and 1 < valor < 100000:
        return date(1899, 12, 30) + timedelta(days=int(valor))  # data em formato numérico do Excel
    conteudo = texto(valor)
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%Y-%m-%d %H:%M:%S",
                    "%d/%m/%Y %H:%M:%S", "%d%m%Y"):
        try:
            return datetime.strptime(conteudo, formato).date()
        except ValueError:
            continue
    raise ErroLinha(f"{nome_campo} inválida: '{conteudo}'. Use o formato DD/MM/AAAA.")


def converter_situacao(valor):
    conteudo = normalizar(texto(valor))
    if not conteudo:
        return None
    if conteudo in ("ativo", "ativa", "a", "sim", "s"):
        return Funcionario.Situacao.ATIVO
    if conteudo in ("inativo", "inativa", "i", "desligado", "desligada", "demitido", "demitida", "nao", "n"):
        return Funcionario.Situacao.INATIVO
    raise ErroLinha(f"Situação '{texto(valor)}' não reconhecida. Use Ativo ou Inativo.")


def chave_contrato(sigla):
    """'TRT 16', 'trt-16' e 'TRT16' são o mesmo contrato."""
    return re.sub(r"[^A-Z0-9]", "", texto(sigla).upper())


# ---------------------------------------------------------------------------
# Importação
# ---------------------------------------------------------------------------
class _Importador:
    def __init__(self, criar_contratos, resultado, contrato_padrao=None):
        self.criar_contratos = criar_contratos
        self.contrato_padrao = contrato_padrao
        self.resultado = resultado
        self.contratos = {chave_contrato(c.sigla): c for c in Contrato.objects.all()}
        self.por_matricula = {f.matricula: f for f in Funcionario.objects.select_related("usuario", "contrato")}
        self.por_cpf = {f.cpf: f for f in self.por_matricula.values()}
        self.linha_da_matricula = {}
        self.linha_do_cpf = {}

    def processar(self, numero, dados):
        matricula = texto(dados.get("matricula"))
        cpf = converter_cpf(dados.get("cpf"))
        if not matricula:
            raise ErroLinha("Matrícula em branco.")
        if not cpf:
            raise ErroLinha("CPF em branco.")
        if matricula in self.linha_da_matricula:
            raise ErroLinha(f"Matrícula repetida na planilha (já aparece na linha {self.linha_da_matricula[matricula]}).")
        if cpf in self.linha_do_cpf:
            raise ErroLinha(f"CPF repetido na planilha (já aparece na linha {self.linha_do_cpf[cpf]}).")
        self.linha_da_matricula[matricula] = numero
        self.linha_do_cpf[cpf] = numero

        existente = self.por_matricula.get(matricula)
        dono_do_cpf = self.por_cpf.get(cpf)
        if dono_do_cpf is not None and dono_do_cpf is not existente:
            raise ErroLinha(f"Este CPF já está cadastrado para outro funcionário "
                            f"({dono_do_cpf.nome_completo}, matrícula {dono_do_cpf.matricula}).")

        valores = {
            "matricula": matricula,
            "cpf": cpf,
            "nome_completo": texto(dados.get("nome_completo")),
            "data_nascimento": converter_data(dados.get("data_nascimento"), "Data de nascimento"),
        }
        if not valores["nome_completo"]:
            raise ErroLinha("Nome em branco.")
        if valores["data_nascimento"] is None:
            raise ErroLinha("Data de nascimento em branco.")
        if not date(1900, 1, 1) < valores["data_nascimento"] < timezone.localdate():
            raise ErroLinha("Data de nascimento fora do esperado. Confira o ano.")

        # Colunas opcionais: só alteram o cadastro se estiverem preenchidas.
        if texto(dados.get("cargo")):
            valores["cargo"] = texto(dados["cargo"])[:120]
        if texto(dados.get("data_admissao")):
            valores["data_admissao"] = converter_data(dados["data_admissao"], "Data de admissão")
        if texto(dados.get("instalacao")):
            valores["instalacao"] = texto(dados["instalacao"])[:200]
        if texto(dados.get("telefone")):
            celular = extrair_celular(dados["telefone"])  # pega o celular mesmo se a célula tiver dois números
            valores["telefone"] = formatar_celular(celular) if celular else texto(dados["telefone"])[:20]
        situacao = converter_situacao(dados.get("situacao"))
        if situacao:
            valores["situacao"] = situacao
        if texto(dados.get("data_desligamento")):
            valores["data_desligamento"] = converter_data(dados["data_desligamento"], "Data de desligamento")

        contrato = self._contrato(dados.get("contrato"), existente)

        if existente:
            funcionario = existente
            alterados = [campo for campo, valor in valores.items() if getattr(funcionario, campo) != valor]
            if funcionario.contrato_id != contrato.pk:
                alterados.append("contrato")
            if not alterados:
                return "sem_alteracao", funcionario, []
            cpf_antigo = funcionario.cpf
        else:
            funcionario = Funcionario()
            alterados = []
            cpf_antigo = None

        for campo, valor in valores.items():
            setattr(funcionario, campo, valor)
        funcionario.contrato = contrato if contrato.pk else None
        funcionario.full_clean(exclude=["contrato"], validate_unique=False)

        if contrato.pk is None:  # contrato novo só é criado se o resto da linha estiver correto
            contrato.save()
            self.resultado.contratos_criados.append(contrato.sigla)
        funcionario.contrato = contrato

        funcionario.save()

        self.por_matricula[matricula] = funcionario
        if cpf_antigo and cpf_antigo != cpf:
            self.por_cpf.pop(cpf_antigo, None)
        self.por_cpf[cpf] = funcionario
        nomes = [str(Funcionario._meta.get_field(c).verbose_name) for c in alterados]
        return ("atualizado" if existente else "criado"), funcionario, nomes

    def _contrato(self, valor, existente):
        sigla = texto(valor).upper()
        if not sigla:
            # Sem contrato na planilha: mantém o atual ou usa o "contrato padrão" escolhido na tela.
            if existente:
                return existente.contrato
            if self.contrato_padrao:
                return self.contrato_padrao
            raise ErroLinha("Contrato em branco. Preencha a coluna Contrato ou escolha um 'contrato padrão' "
                            "na tela de importação.")
        contrato = self.contratos.get(chave_contrato(sigla)) or self._contrato_pelo_inicio(sigla)
        if contrato is None:
            if not self.criar_contratos:
                raise ErroLinha(f"Contrato '{sigla}' não está cadastrado. Cadastre em 'Contratos / lotações' "
                                "ou marque a opção de criar contratos automaticamente.")
            contrato = Contrato(sigla=sigla[:30])
            self.contratos[chave_contrato(sigla)] = contrato
        return contrato

    def _contrato_pelo_inicio(self, texto_planilha):
        """
        'SEMED - CT 11/2024 FERISTA' -> contrato SEMED (o texto começa com a sigla).
        Se mais de uma sigla servir, vale a mais longa. A associação aparece no relatório.
        """
        chave = chave_contrato(texto_planilha)
        candidatos = [(len(c), c) for c, contrato in self.contratos.items() if contrato.pk and c and chave.startswith(c)]
        if not candidatos:
            return None
        contrato = self.contratos[max(candidatos)[1]]
        self.resultado.contratos_associados[texto_planilha] = contrato.sigla
        return contrato


def _mensagem_validacao(erro):
    if hasattr(erro, "message_dict"):
        partes = []
        for campo, mensagens in erro.message_dict.items():
            try:
                rotulo = str(Funcionario._meta.get_field(campo).verbose_name).capitalize()
            except Exception:
                rotulo = ""
            partes.append((f"{rotulo}: " if rotulo else "") + " ".join(mensagens))
        return " | ".join(partes)
    return " ".join(erro.messages)


def importar_funcionarios(arquivo, nome_arquivo, criar_contratos=False, simulacao=False, contrato_padrao=None):
    resultado = ResultadoImportacao(simulacao=simulacao)
    try:
        linhas = ler_planilha(arquivo, nome_arquivo)
        posicao_cabecalho, mapa = localizar_cabecalho(linhas)
    except ErroPlanilha as erro:
        resultado.erro_geral = str(erro)
        return resultado

    resultado.colunas_reconhecidas = [COLUNAS[campo][0] for campo in mapa.values()]
    linhas_dados = linhas[posicao_cabecalho + 1:]
    if len(linhas_dados) > LIMITE_LINHAS:
        resultado.erro_geral = f"A planilha tem mais de {LIMITE_LINHAS} linhas. Divida em arquivos menores."
        return resultado

    with transaction.atomic():
        importador = _Importador(criar_contratos, resultado, contrato_padrao)
        for numero, valores in linhas_dados:
            dados = {campo: (valores[coluna] if coluna < len(valores) else None) for coluna, campo in mapa.items()}
            if all(texto(v) == "" for v in dados.values()):
                continue  # linha vazia
            resultado.total_linhas += 1
            identificacao = " — ".join(filter(None, [texto(dados.get("matricula")), texto(dados.get("nome_completo"))])) or "—"
            try:
                with transaction.atomic():
                    situacao, funcionario, alterados = importador.processar(numero, dados)
            except ErroLinha as erro:
                resultado.erros.append((numero, identificacao, str(erro)))
            except ValidationError as erro:
                resultado.erros.append((numero, identificacao, _mensagem_validacao(erro)))
            except IntegrityError:
                resultado.erros.append((numero, identificacao, "Matrícula ou CPF já cadastrado para outro funcionário."))
            else:
                descricao = f"{funcionario.nome_completo} ({funcionario.matricula})"
                if situacao == "criado":
                    resultado.criados.append((numero, descricao))
                elif situacao == "atualizado":
                    resultado.atualizados.append((numero, descricao, ", ".join(alterados)))
                else:
                    resultado.sem_alteracao += 1
        if simulacao:
            transaction.set_rollback(True)  # desfaz tudo: foi só uma conferência
    return resultado


# ---------------------------------------------------------------------------
# Modelo de planilha para download
# ---------------------------------------------------------------------------
def gerar_modelo_planilha():
    livro = Workbook()
    aba = livro.active
    aba.title = "Funcionários"
    titulos = [titulo for titulo, _, _ in COLUNAS.values()]
    aba.append(titulos)

    obrigatoria = PatternFill("solid", fgColor="1F4E79")
    opcional = PatternFill("solid", fgColor="7F7F7F")
    for coluna, (campo, (titulo, eh_obrigatoria, _)) in enumerate(COLUNAS.items(), start=1):
        celula = aba.cell(row=1, column=coluna)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = obrigatoria if eh_obrigatoria else opcional
        celula.alignment = Alignment(horizontal="center")
        letra = celula.column_letter
        aba.column_dimensions[letra].width = 40 if campo == "nome_completo" else 20
        formato = "DD/MM/YYYY" if campo.startswith("data_") else "@" if campo in ("matricula", "cpf", "telefone") else None
        if formato:
            for linha in range(2, 1002):
                aba.cell(row=linha, column=coluna).number_format = formato
    aba.freeze_panes = "A2"

    posicao = {campo: i for i, campo in enumerate(COLUNAS, start=1)}
    situacao = DataValidation(type="list", formula1='"Ativo,Inativo"', allow_blank=True)
    aba.add_data_validation(situacao)
    situacao.add(f"{aba.cell(row=2, column=posicao['situacao']).column_letter}2:"
                 f"{aba.cell(row=2, column=posicao['situacao']).column_letter}1001")
    siglas = ",".join(Contrato.objects.filter(ativo=True).values_list("sigla", flat=True))
    if siglas and len(siglas) < 250:
        lista_contratos = DataValidation(type="list", formula1=f'"{siglas}"', allow_blank=True)
        aba.add_data_validation(lista_contratos)
        letra = aba.cell(row=2, column=posicao["contrato"]).column_letter
        lista_contratos.add(f"{letra}2:{letra}1001")

    instrucoes = livro.create_sheet("Instruções")
    instrucoes.column_dimensions["A"].width = 110
    textos = [
        "COMO PREENCHER A PLANILHA DE FUNCIONÁRIOS",
        "",
        "1. Preencha a aba 'Funcionários', um funcionário por linha, a partir da linha 2. Não altere os títulos.",
        "2. Colunas em AZUL são obrigatórias: " + ", ".join(COLUNAS[c][0] for c in CAMPOS_OBRIGATORIOS) + ".",
        "3. Colunas em CINZA são opcionais. Se ficarem em branco, o cadastro existente não é alterado.",
        "4. CPF: pode ser com ou sem pontuação (000.000.000-00 ou 00000000000).",
        "5. Datas no formato DD/MM/AAAA (ex.: 07/03/1990).",
        "6. Contrato: use a sigla cadastrada no sistema (ex.: SEMED, SEMCAS, TRT-16, TJMA). Se deixar em branco, "
        "o sistema usa o 'contrato padrão' escolhido na tela de importação.",
        "7. Situação: Ativo ou Inativo. Em branco = Ativo para funcionários novos.",
        "8. Se a matrícula já existir no sistema, os dados do funcionário serão ATUALIZADOS.",
        "9. Telefone: o CELULAR com DDD (ex.: 98 98888-7777). É por ele que o assistente do WhatsApp reconhece a pessoa.",
        "",
        "Dica: use a opção 'Somente conferir' no sistema para ver os erros antes de importar de verdade.",
    ]
    for texto_linha in textos:
        instrucoes.append([texto_linha])
    instrucoes["A1"].font = Font(bold=True, size=14)

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
