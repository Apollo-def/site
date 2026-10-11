"""Gera planilhas Excel (.xlsx) para download no painel do RH (exportações)."""
import io
from datetime import date, datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _celula(valor):
    if isinstance(valor, datetime):
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    return valor


def resposta_excel(nome_arquivo, titulo_aba, titulos, linhas):
    """
    titulos: nomes das colunas; linhas: listas de valores (datas viram datas de verdade no Excel).
    Devolve o HttpResponse com o arquivo para baixar.
    """
    return resposta_excel_abas(nome_arquivo, [(titulo_aba, titulos, linhas)])


def resposta_excel_abas(nome_arquivo, abas):
    """Uma planilha com várias abas: abas = [(título da aba, títulos das colunas, linhas)]."""
    livro = Workbook()
    livro.remove(livro.active)
    for titulo_aba, titulos, linhas in abas:
        _preencher_aba(livro.create_sheet(titulo_aba[:31]), titulos, linhas)
    saida = io.BytesIO()
    livro.save(saida)
    resposta = HttpResponse(saida.getvalue(), content_type=TIPO_XLSX)
    resposta["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'
    resposta["Cache-Control"] = "private, no-store"
    return resposta


def _preencher_aba(aba, titulos, linhas):
    aba.append(titulos)
    for coluna in range(1, len(titulos) + 1):
        celula = aba.cell(row=1, column=coluna)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0B3D6E")
        celula.alignment = Alignment(vertical="center")
    larguras = [len(str(t)) for t in titulos]
    for linha in linhas:
        valores = [_celula(v) for v in linha]
        aba.append(valores)
        for i, valor in enumerate(valores):
            larguras[i] = max(larguras[i], min(len(str(valor or "")), 60))
    for i, largura in enumerate(larguras, start=1):
        aba.column_dimensions[get_column_letter(i)].width = largura + 3
        for celula in aba[get_column_letter(i)][1:]:
            if isinstance(celula.value, datetime):
                celula.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(celula.value, date):
                celula.number_format = "DD/MM/YYYY"
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
