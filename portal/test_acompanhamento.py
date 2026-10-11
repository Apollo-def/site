import io
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.base import ContentFile
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from auditoria.models import LogAuditoria
from auditoria.servicos import registrar
from contracheques.models import Contracheque
from contracheques.tests import Base, gerar_pdf, pagina_contracheque
from solicitacoes.models import Solicitacao
from solicitacoes.servicos import criar_solicitacao_vt, tipos_vt_disponiveis
from whatsapp.models import Conversa

PAINEL = reverse("painel_acompanhamento")


def ler_excel(resposta):
    aba = load_workbook(io.BytesIO(resposta.content)).active
    return [list(linha) for linha in aba.iter_rows(values_only=True)]


class PainelAcompanhamentoTestes(Base):
    def setUp(self):
        super().setUp()
        self.enviar(gerar_pdf([pagina_contracheque("7013", "CLARA MARIA SILVA SANTOS"),
                               pagina_contracheque("6068", "JOAO PEREIRA")]))
        # Clara conversou com o assistente e recebeu o contracheque; João ainda não.
        Conversa.objects.create(numero="5598987654321", funcionario=self.clara)
        registrar(LogAuditoria.Acao.WHATSAPP_DOCUMENTO_ENVIADO, objeto=Contracheque.objects.get(funcionario=self.clara),
                  identificador=self.clara.matricula)
        self.client.force_login(self.rh)

    def test_mostra_quem_recebeu_e_quem_nao_recebeu(self):
        resposta = self.client.get(PAINEL)
        self.assertContains(resposta, "Contracheque de Janeiro de 2026")
        self.assertContains(resposta, "1 recebidos (50%)")
        self.assertContains(resposta, "Ainda não receberam (1)")
        self.assertContains(resposta, "JOAO PEREIRA")
        self.assertEqual(resposta.context["pct_entraram"], 50)  # Clara já conversou; João nunca

    def test_exporta_quem_nao_abriu_em_excel_e_registra(self):
        resposta = self.client.get(PAINEL, {"competencia": "2026-01", "exportar": "nao_abriram"})
        self.assertEqual(resposta["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        linhas = ler_excel(resposta)
        self.assertEqual(linhas[0][:3], ["Nome", "Matrícula", "CPF"])
        self.assertEqual([l[0] for l in linhas[1:]], ["JOAO PEREIRA"])
        self.assertTrue(LogAuditoria.objects.filter(acao="exportou_dados").exists())

    def test_pedido_parado_aparece(self):
        tipo = tipos_vt_disponiveis(self.joao)["vt_utilizar"]
        pedido = Solicitacao.objects.create(funcionario=self.joao, tipo=tipo)
        Solicitacao.objects.filter(pk=pedido.pk).update(atualizado_em=timezone.now() - timedelta(days=5))
        resposta = self.client.get(PAINEL)
        self.assertContains(resposta, "Pedidos parados há mais de 3 dias (1)")
        self.assertContains(resposta, pedido.protocolo)

    def test_atalho_na_tela_inicial_do_painel(self):
        self.assertContains(self.client.get(reverse("admin:index")), PAINEL)

    def test_quem_nao_pode_ver_funcionarios_nao_entra(self):
        limitado = get_user_model().objects.create_user("limitado", password="SenhaForte!2026", is_staff=True)
        self.client.force_login(limitado)
        self.assertEqual(self.client.get(PAINEL).status_code, 403)


class ExportarSolicitacoesTestes(Base):
    def setUp(self):
        super().setUp()
        tipos = tipos_vt_disponiveis(self.clara)
        self.pedido = Solicitacao.objects.create(funcionario=self.clara, tipo=tipos["vt_cancelar"])
        Solicitacao.objects.create(funcionario=self.joao, tipo=tipos["vt_utilizar"], status="concluida",
                                   resposta_rh="Incluído")
        self.client.force_login(self.rh)
        self.url = reverse("admin:solicitacoes_solicitacao_exportar")

    def test_exporta_tudo_ou_so_o_filtrado(self):
        tudo = ler_excel(self.client.get(self.url))
        self.assertEqual(len(tudo) - 1, 2)
        self.assertIn("VT", tudo[0])
        so_pendentes = ler_excel(self.client.get(self.url, {"status__exact": "pendente"}))
        self.assertEqual(len(so_pendentes) - 1, 1)
        self.assertEqual(so_pendentes[1][0], self.pedido.protocolo)
        self.assertEqual(so_pendentes[1][8], "Cancelar")

    def test_botao_na_lista(self):
        self.assertContains(self.client.get(reverse("admin:solicitacoes_solicitacao_changelist")), self.url)


@override_settings(EMAIL_HOST="smtp.teste", EMAILS_AVISO_RH=["rh@empresa.com.br"])
class AvisoPorEmailTestes(Base):
    def devolver_termo(self):
        tipo = tipos_vt_disponiveis(self.clara)["vt_utilizar"]
        with self.captureOnCommitCallbacks(execute=True):
            return criar_solicitacao_vt(self.clara, tipo, ContentFile(b"%PDF-1.4 termo"), ".pdf",
                                        origem=Solicitacao.Origem.WHATSAPP)

    def test_rh_recebe_email_sem_cpf(self):
        pedido = self.devolver_termo()
        self.assertEqual(len(mail.outbox), 1)
        email = mail.outbox[0]
        self.assertEqual(email.to, ["rh@empresa.com.br"])
        self.assertIn(pedido.protocolo, email.subject)
        self.assertIn("CLARA MARIA SILVA SANTOS", email.body)
        self.assertIn(reverse("admin:solicitacoes_solicitacao_change", args=[pedido.pk]), email.body)
        self.assertNotIn(self.clara.cpf, email.body)

    @override_settings(EMAIL_HOST="")
    def test_sem_configuracao_nao_envia_e_o_pedido_funciona(self):
        self.devolver_termo()
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(Solicitacao.objects.count(), 1)
