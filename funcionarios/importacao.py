"""
Importação de funcionários por planilha Excel (.xlsx).

- O funcionário é identificado pela MATRÍCULA: se já existir, é atualizado; se não, é cadastrado.
- Obrigatórias: Nome completo, Matrícula e CPF. As demais colunas são opcionais e, se vierem em branco,
  não apagam o que já está cadastrado.
- Cada linha é independente: uma linha com erro não impede as outras.
- No modo "somente conferir" tudo é processado e depois desfeito (nada fica salvo).
"""
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from .models import Funcionario, Instalacao, somente_digitos

# campo: (título principal, obrigatório?, outros títulos aceitos)
COLUNAS = {
    "nome_completo": ("Nome completo", True, ["nome", "funcionario", "colaborador", "empregado"]),
    "matricula": ("Matrícula", True, ["mat", "matr", "registro"]),
    "cpf": ("CPF", True, []),
    "data_nascimento": ("Data de nascimento", False, ["nascimento", "dtnasc", "datanasc", "dtnascimento"]),
    "cargo": ("Cargo", False, ["funcao", "tipodeposto"]),
    "instalacao": ("Instalação", False, ["localdetrabalho", "unidade", "lotacao"]),
    "escala": ("Escala", False, []),
    "telefone": ("Telefone", False, ["celular", "whatsapp", "fone"]),
    "situacao": ("Situação", False, ["status"]),
}
OBRIGATORIOS = [c for c, (_, obrigatorio, _) in COLUNAS.items() if obrigatorio]


class ErroPlanilha(Exception):
    pass


class ErroLinha(Exception):
    pass


@dataclass
class Resultado:
    simulacao: bool = False
    total: int = 0
    criados: int = 0
    atualizados: int = 0
    sem_alteracao: int = 0
    erros: list = field(default_factory=list)  # [(linha, identificação, mensagem)]
    instalacoes_criadas: list = field(default_factory=list)
    erro_geral: str = ""

    def resumo(self):
        texto = (f"{self.total} linhas: {self.criados} cadastrados, {self.atualizados} atualizados, "
                 f"{self.sem_alteracao} sem alteração, {len(self.erros)} com erro")
        return texto + (" (somente conferência — nada foi salvo)" if self.simulacao else "")


def normalizar(valor):
    valor = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", valor.lower())


def corrigir_acentos(valor):
    """'SAO LUÃ\\x8dS' (acento quebrado por conversão errada de arquivo) -> 'SAO LUÍS'."""
    if "Ã" in valor or "Â" in valor:
        for codificacao in ("cp1252", "latin-1"):
            try:
                return valor.encode(codificacao).decode("utf-8")
            except (UnicodeEncodeError, UnicodeDecodeError):
                continue
    return valor


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return corrigir_acentos(" ".join(str(valor).split()))


def converter_cpf(valor):
    cpf = somente_digitos(texto(valor))
    return cpf.zfill(11) if 0 < len(cpf) < 11 else cpf  # o Excel apaga os zeros à esquerda


def converter_data(valor):
    if valor is None or texto(valor) == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and 1 < valor < 100000:
        return date(1899, 12, 30) + timedelta(days=int(valor))
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto(valor), formato).date()
        except ValueError:
            continue
    raise ErroLinha(f"Data de nascimento inválida: '{texto(valor)}'. Use DD/MM/AAAA.")


def converter_situacao(valor):
    conteudo = normalizar(valor)
    if not conteudo:
        return None
    if conteudo in ("ativo", "ativa", "a"):
        return Funcionario.Situacao.ATIVO
    if conteudo in ("inativo", "inativa", "i", "desligado", "desligada"):
        return Funcionario.Situacao.INATIVO
    raise ErroLinha(f"Situação '{texto(valor)}' não reconhecida. Use Ativo ou Inativo.")


def ler_linhas(arquivo):
    try:
        aba = load_workbook(arquivo, read_only=True, data_only=True).worksheets[0]
    except Exception as erro:
        raise ErroPlanilha("Não foi possível abrir o arquivo. Envie uma planilha Excel (.xlsx).") from erro
    return [(numero, list(valores)) for numero, valores in enumerate(aba.iter_rows(values_only=True), 1)]


def localizar_cabecalho(linhas):
    aceitos = {}
    for campo, (titulo, _, apelidos) in COLUNAS.items():
        for nome in (titulo, campo, *apelidos):
            aceitos[normalizar(nome)] = campo
    for posicao, (_, valores) in enumerate(linhas[:15]):
        mapa = {}
        for coluna, valor in enumerate(valores):
            campo = aceitos.get(normalizar(valor))
            if campo and campo not in mapa.values():
                mapa[coluna] = campo
        if all(c in mapa.values() for c in OBRIGATORIOS):
            return posicao, mapa
    raise ErroPlanilha("Não encontrei a linha de títulos. A planilha precisa ter as colunas "
                       + ", ".join(COLUNAS[c][0] for c in OBRIGATORIOS) + ".")


class _Importador:
    def __init__(self, resultado):
        self.resultado = resultado
        self.instalacoes = {normalizar(i.nome): i for i in Instalacao.objects.all()}
        self.por_matricula = {f.matricula: f for f in Funcionario.objects.all()}
        self.por_cpf = {f.cpf: f for f in self.por_matricula.values()}
        self.vistos_matricula = {}
        self.vistos_cpf = {}

    def processar(self, numero, dados):
        matricula = texto(dados.get("matricula"))
        cpf = converter_cpf(dados.get("cpf"))
        nome = texto(dados.get("nome_completo"))
        if not matricula:
            raise ErroLinha("Matrícula em branco.")
        if not cpf:
            raise ErroLinha("CPF em branco.")
        if not nome:
            raise ErroLinha("Nome em branco.")
        if matricula in self.vistos_matricula:
            raise ErroLinha(f"Matrícula repetida (já aparece na linha {self.vistos_matricula[matricula]}).")
        if cpf in self.vistos_cpf:
            raise ErroLinha(f"CPF repetido (já aparece na linha {self.vistos_cpf[cpf]}).")
        self.vistos_matricula[matricula] = numero
        self.vistos_cpf[cpf] = numero

        existente = self.por_matricula.get(matricula)
        dono_cpf = self.por_cpf.get(cpf)
        if dono_cpf is not None and dono_cpf is not existente:
            raise ErroLinha(f"CPF já cadastrado para {dono_cpf}.")

        valores = {"matricula": matricula, "cpf": cpf, "nome_completo": nome}
        for campo in ("cargo", "escala", "telefone"):
            if texto(dados.get(campo)):
                valores[campo] = texto(dados[campo])[:Funcionario._meta.get_field(campo).max_length]
        if texto(dados.get("data_nascimento")):
            valores["data_nascimento"] = converter_data(dados["data_nascimento"])
        situacao = converter_situacao(dados.get("situacao"))
        if situacao:
            valores["situacao"] = situacao
        nome_instalacao = texto(dados.get("instalacao")).upper()[:200]
        if nome_instalacao:
            valores["instalacao"] = self._instalacao(nome_instalacao)

        funcionario = existente or Funcionario()
        if existente and all(getattr(existente, c) == v for c, v in valores.items()):
            return "sem_alteracao"
        for campo, valor in valores.items():
            setattr(funcionario, campo, valor)
        funcionario.full_clean(exclude=["instalacao"], validate_unique=False)
        # A instalação nova só é gravada depois que o resto da linha passou na validação.
        if funcionario.instalacao and funcionario.instalacao.pk is None:
            funcionario.instalacao.save()
            self.resultado.instalacoes_criadas.append(funcionario.instalacao.nome)
        funcionario.instalacao = funcionario.instalacao  # atualiza o id depois de salvar
        funcionario.save()
        self.por_matricula[matricula] = funcionario
        self.por_cpf[cpf] = funcionario
        return "atualizado" if existente else "criado"

    def _instalacao(self, nome):
        chave = normalizar(nome)
        if chave not in self.instalacoes:
            self.instalacoes[chave] = Instalacao(nome=nome)  # só é salva se a linha estiver correta
        return self.instalacoes[chave]


def importar_funcionarios(arquivo, simulacao=False):
    resultado = Resultado(simulacao=simulacao)
    try:
        linhas = ler_linhas(arquivo)
        posicao, mapa = localizar_cabecalho(linhas)
    except ErroPlanilha as erro:
        resultado.erro_geral = str(erro)
        return resultado

    with transaction.atomic():
        importador = _Importador(resultado)
        for numero, valores in linhas[posicao + 1:]:
            dados = {campo: (valores[col] if col < len(valores) else None) for col, campo in mapa.items()}
            if all(texto(v) == "" for v in dados.values()):
                continue
            resultado.total += 1
            identificacao = " — ".join(filter(None, [texto(dados.get("matricula")),
                                                     texto(dados.get("nome_completo"))])) or "—"
            try:
                with transaction.atomic():
                    situacao = importador.processar(numero, dados)
            except ErroLinha as erro:
                resultado.erros.append((numero, identificacao, str(erro)))
            except ValidationError as erro:
                resultado.erros.append((numero, identificacao, " ".join(erro.messages)))
            except IntegrityError:
                resultado.erros.append((numero, identificacao, "Matrícula ou CPF já cadastrado."))
            else:
                if situacao == "criado":
                    resultado.criados += 1
                elif situacao == "atualizado":
                    resultado.atualizados += 1
                else:
                    resultado.sem_alteracao += 1
        if simulacao:
            transaction.set_rollback(True)
    return resultado
