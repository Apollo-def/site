import base64
import io
import json
from datetime import date, datetime

from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from contas.models import Contrato, Funcionario
from contracheques.tests import Base

from .leitura import encontrar_funcionario, ler_mensagem
from .models import Ocorrencia

TOKEN = "senha-do-gateway-com-mais-de-20"
MENSAGEM_REAL = ("Extra hoje 08/10\nFuncionária: Rosalinda Silva Pereira.\n"
                 "Cobrindo Atestado Médico: Gercileia Sousa Fontes.\nEscola: lindalva Teotônio")
QUANDO = int(datetime(2026, 10, 8, 11, 4, tzinfo=timezone.get_current_timezone()).timestamp())


@override_settings(WHATSAPP_CANAL="gateway", WHATSAPP_GATEWAY_TOKEN=TOKEN)
class OcorrenciasTestes(Base):
    def setUp(self):
        super().setUp()
        semed = Contrato.objects.get(sigla="SEMED")

        def criar(nome, matricula, cpf, instalacao):
            return Funcionario.objects.create(nome_completo=nome, matricula=matricula, cpf=cpf, contrato=semed,
                                              data_nascimento=date(1990, 1, 1), instalacao=instalacao)

        self.rosalinda = criar("ROSALINDA SILVA PEREIRA", "6908", "39053344705",
                              "ITAQUI - BACANGA - LINDALVA TEOTONIA NUNES- FUNDAMENTAL")
        self.gercilea = criar("GERCILEA SOUSA FONTES", "6576", "71428793860", "FERISTAS DEPTO 162")
        self.maria_a = criar("MARIA JOSE SILVA", "1001", "86288366757", "CENTRO - BERNARDINA SPÃ\x8dNDOLA- INFANTIL")
        self.maria_b = criar("MARIA JOSE DA SILVA COSTA", "1002", "52184383010", "ANIL - ESPERANCA DO AMANHA")

    def postar(self, texto, id_msg="m1", midia=None, autor="5598911110000"):
        dados = {"tipo": "ocorrencia_grupo", "grupo": "1203@g.us", "id": id_msg, "texto": texto, "midia": midia,
                 "autor": autor, "autor_nome": "Wenderson Freitas - SEMED", "enviada_em": QUANDO}
        return self.client.post("/whatsapp/gateway/", json.dumps(dados), content_type="application/json",
                                headers={"X-Gateway-Token": TOKEN})

    def test_mensagem_real_do_grupo_vira_ocorrencia_conferida(self):
        self.assertEqual(self.postar(MENSAGEM_REAL).status_code, 200)
        o = Ocorrencia.objects.get()
        self.assertEqual((o.tipo, o.data, o.funcionario, o.coberto, o.situacao),
                         ("extra", date(2026, 10, 8), self.rosalinda, self.gercilea, "ok"))  # "Gercileia" ~ GERCILEA
        self.assertEqual(o.observacao, "Atestado Médico")
        self.assertEqual(o.supervisor_nome, "Wenderson Freitas - SEMED")
        self.postar(MENSAGEM_REAL)  # a mesma mensagem de novo: não duplica
        self.assertEqual(Ocorrencia.objects.count(), 1)

    def test_modelo_nome_em_outra_ordem_e_escola_desempata(self):
        self.postar("*TIPO:* FALTA\n*DATA:* 07/10/2026\n*FUNCIONÁRIO:* Silva Maria José\n"
                    "*ESCOLA/POSTO:* Bernardina Spíndola")  # duas "Maria José Silva": a escola decide
        o = Ocorrencia.objects.get()
        self.assertEqual((o.tipo, o.data, o.funcionario, o.situacao), ("falta", date(2026, 10, 7), self.maria_a, "ok"))

    def test_na_duvida_vai_para_o_rh_conferir(self):
        self.postar("*TIPO:* FALTA / ATESTADO / COBERTURA / EXTRA\n*FUNCIONÁRIO:* Maria José Silva", "m1")
        o = Ocorrencia.objects.get()
        self.assertIsNone(o.funcionario)  # duas pessoas possíveis e sem escola: não chuta
        self.assertEqual(o.situacao, "conferir")
        self.assertIn("2 pessoas", o.motivo_conferir)
        self.assertIn("tipo", o.motivo_conferir)  # modelo colado sem escolher o tipo
        self.assertIn("data", o.motivo_conferir)

    def test_conversa_do_grupo_e_ignorada_e_foto_junta_na_ocorrencia(self):
        self.postar("Bom dia pessoal! Alguém sabe se tem reunião hoje?", "c1")
        self.assertFalse(Ocorrencia.objects.exists())
        self.postar("Atestado hoje\nFuncionária: Rosalinda Pereira\nEscola: Lindalva", "a1")
        foto = {"base64": base64.b64encode(b"\xff\xd8\xff\xe0foto do atestado").decode(), "nome": "foto.jpg"}
        self.postar("", "a2", midia=foto)  # a foto do atestado veio logo depois, sem texto
        o = Ocorrencia.objects.get()
        self.assertEqual((o.tipo, o.funcionario), ("atestado", self.rosalinda))
        self.assertTrue(o.anexo.name.startswith("ocorrencias/"))

    def test_leitura_de_datas_e_periodo(self):
        lido = ler_mensagem("*TIPO:* ATESTADO\n*DATA:* 08/10\n*FUNCIONÁRIO:* Rosalinda Pereira\n"
                            "*ATESTADO:* de 08/10 a 10/10", timezone.now().replace(year=2026, month=10, day=8))
        self.assertEqual((lido["atestado_inicio"], lido["atestado_fim"]), (date(2026, 10, 8), date(2026, 10, 10)))
        self.assertEqual(encontrar_funcionario("Rosalinda", "")[0], None)  # só o primeiro nome não basta
        self.assertEqual(encontrar_funcionario("Fulano", "", "6908")[0], self.rosalinda)  # matrícula manda

    def test_painel_resumo_conferencia_e_planilha(self):
        self.postar(MENSAGEM_REAL, "m1")
        self.postar("Falta ontem\nFuncionária: Maria José Silva", "m2")
        self.client.force_login(self.rh)
        lista = self.client.get(reverse("admin:ocorrencias_ocorrencia_changelist"))
        self.assertContains(lista, "Para o RH conferir")
        self.assertContains(lista, "ROSALINDA SILVA PEREIRA")
        duvida = Ocorrencia.objects.get(situacao="conferir")
        tela = self.client.get(reverse("admin:ocorrencias_ocorrencia_change", args=[duvida.pk]))
        self.assertContains(tela, "MARIA JOSE DA SILVA COSTA")  # candidatos para o RH escolher
        self.client.post(reverse("admin:ocorrencias_ocorrencia_change", args=[duvida.pk]), {
            "tipo": "falta", "data": "07/10/2026", "funcionario": self.maria_b.pk, "coberto": "",
            "atestado_inicio": "", "atestado_fim": "", "observacao": "", "situacao": "conferir"})
        duvida.refresh_from_db()
        self.assertEqual((duvida.funcionario, duvida.situacao, duvida.conferida_por), (self.maria_b, "ok", self.rh))
        planilha = self.client.get(reverse("admin:ocorrencias_ocorrencia_exportar") + "?tipo__exact=extra")
        linhas = list(load_workbook(io.BytesIO(planilha.content)).active.values)
        self.assertEqual(len(linhas), 2)
        self.assertEqual(linhas[1][1:4], ("Extra", "ROSALINDA SILVA PEREIRA", "6908"))

    def test_relatorio_detalhado_do_periodo(self):
        self.postar(MENSAGEM_REAL, "r1")  # extra da Rosalinda cobrindo a Gercilea
        self.postar("*TIPO:* ATESTADO\n*DATA:* 06/10/2026\n*FUNCIONÁRIO:* Gercilea Fontes\n"
                    "*ATESTADO:* de 06/10 a 08/10", "r2")
        self.postar("*TIPO:* FALTA\n*DATA:* 07/10/2026\n*FUNCIONÁRIO:* Rosalinda Pereira\n*ESCOLA/POSTO:* Lindalva", "r3")
        self.postar("*TIPO:* FALTA\n*DATA:* 15/09/2026\n*FUNCIONÁRIO:* Rosalinda Pereira", "fora")  # fora do período
        self.client.force_login(self.rh)
        url = reverse("admin:ocorrencias_ocorrencia_relatorio") + "?de=2026-10-01&ate=2026-10-31"
        tela = self.client.get(url)
        self.assertEqual(tela.status_code, 200)
        dados = tela.context["dados"]
        self.assertEqual({k: dados["totais"][k] for k in ("total", "falta", "atestado", "extra", "dias_atestado")},
                         {"total": 3, "falta": 1, "atestado": 1, "extra": 1, "dias_atestado": 3})
        rosalinda = next(p for p in dados["por_funcionario"] if p["matricula"] == "6908")
        gercilea = next(p for p in dados["por_funcionario"] if p["matricula"] == "6576")
        self.assertEqual((rosalinda["falta"], rosalinda["extra"], rosalinda["total"]), (1, 1, 2))
        self.assertEqual((gercilea["atestado"], gercilea["dias_atestado"], gercilea["coberto"]), (1, 3, 1))
        self.assertEqual([d["data"] for d in dados["por_dia"]], [date(2026, 10, 6), date(2026, 10, 7), date(2026, 10, 8)])
        self.assertContains(tela, "Imprimir / salvar PDF")
        planilha = self.client.get(url + "&formato=excel")
        livro = load_workbook(io.BytesIO(planilha.content))
        self.assertEqual(livro.sheetnames, ["Detalhado", "Por funcionário", "Por escola", "Por dia", "Por supervisor"])
        self.assertEqual(livro["Detalhado"].max_row, 4)  # título + 3 ocorrências do período
        so_faltas = self.client.get(url + "&tipo=falta").context["dados"]
        self.assertEqual(so_faltas["totais"]["total"], 1)

    def reagir(self, id_reagido, emoji="👍", de_mim=True, id_reacao="re1"):
        dados = {"tipo": "reacao_grupo", "grupo": "1203@g.us", "id": id_reacao, "id_reagido": id_reagido,
                 "emoji": emoji, "de_mim": de_mim, "autor_nome": "ADM"}
        return self.client.post("/whatsapp/gateway/", json.dumps(dados), content_type="application/json",
                                headers={"X-Gateway-Token": TOKEN})

    def test_joinha_do_adm_marca_ok_e_tirar_desmarca(self):
        self.postar("Falta\n08/10.\nColaborador: Rosalinda Silva Pereira\nPosto: Lindalva\n\nSem cobertura", "f1")
        o = Ocorrencia.objects.get()
        self.assertEqual((o.tipo, o.observacao), ("falta", "Sem cobertura"))  # "sem cobertura" não é cobertura
        self.reagir("f1", "👍", de_mim=False)  # 👍 de outra pessoa do grupo não conta
        o.refresh_from_db()
        self.assertIsNone(o.ok_adm_em)
        self.reagir("f1", "👍🏽")  # o ADM (celular do chip), com cor de pele
        o.refresh_from_db()
        self.assertIsNotNone(o.ok_adm_em)
        self.client.force_login(self.rh)
        lista = self.client.get(reverse("admin:ocorrencias_ocorrencia_changelist") + "?ok_adm=sim")
        self.assertContains(lista, "ROSALINDA SILVA PEREIRA")
        self.assertNotContains(self.client.get(reverse("admin:ocorrencias_ocorrencia_changelist") + "?ok_adm=nao"),
                               "ROSALINDA SILVA PEREIRA")
        self.reagir("f1", "")  # tirou o 👍
        o.refresh_from_db()
        self.assertIsNone(o.ok_adm_em)

    def test_varias_ocorrencias_numa_mensagem_e_robo_so_quando_registrou(self):
        semed = Contrato.objects.get(sigla="SEMED")
        for nome, matricula, cpf in (("ROGERIO CICERO ALMEIDA VIANA", "5221", "11122233396"),
                                     ("IVAN DANGELO CARVALHO", "5335", "22233344405"),
                                     ("OTAVIO FERNANDES COSTA FILHO", "6420", "33344455514")):
            Funcionario.objects.create(nome_completo=nome, matricula=matricula, cpf=cpf, contrato=semed,
                                       data_nascimento=date(1990, 1, 1))
        relato = ("Boa tarde!\nFalta\nDia 08/10\nFuncionário Rogério Cícero Almeida\nPosto Dom José Medeiros Delgado\n"
                  "OBS o funcionário está cobrindo férias do funcionário\nJonas Pereira lima\n\n"
                  "Falta\nDia 08/10\nFuncionário Ivan D\nÂngelo carvalho\nPosto UEB Rio grande infantil\n"
                  "Sem cobertura\n\nExtra\nDia 08/10\nFuncionário Otávio Fernandes Costa\nPosto Haydée chaves\n"
                  "O mesmo está cobrindo Aviso do Rogério Cícero Almeida")
        resposta = self.postar(relato, "relato")
        self.assertEqual(resposta.json(), {"reagir": "🤖"})
        ocorrencias = list(Ocorrencia.objects.order_by("id_externo"))
        self.assertEqual([o.id_externo for o in ocorrencias], ["relato#1", "relato#2", "relato#3"])
        self.assertEqual([(o.tipo, o.funcionario.matricula, o.situacao) for o in ocorrencias],
                         [("falta", "5221", "ok"), ("falta", "5335", "ok"), ("extra", "6420", "ok")])
        self.assertEqual((ocorrencias[1].observacao, ocorrencias[2].coberto.matricula, ocorrencias[2].observacao),
                         ("Sem cobertura", "5221", "Aviso"))
        self.assertEqual(self.postar(relato, "relato").json(), {})  # repetida: não duplica nem reage de novo
        self.assertEqual(self.postar("Boa tarde pessoal!", "oi").json(), {})  # conversa: sem 🤖
        self.assertEqual(Ocorrencia.objects.count(), 3)
